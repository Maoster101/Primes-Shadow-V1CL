"""API routes — wiring hub + small endpoints (health, models, settings, upload, events).

Sub-routers handle the heavy lifting:
  chats.py     — chat CRUD + send_message streaming
  sessions.py  — frame endpoints + 3D graph builder
  drafts.py    — draft lifecycle, library/tentative, dream-all, extract-proposals
  corpus_routes.py — corpus browser, node CRUD, collections
  mining.py    — conversation + narrative mining, push-mined, scratch sessions
"""
from __future__ import annotations
import logging

logger = logging.getLogger(__name__)
from fastapi import APIRouter, HTTPException, UploadFile, File
from pydantic import BaseModel

from ..services import ollama
from . import deps
from .deps import chat_store, event_log

# Sub-routers
from . import chats, sessions, drafts, corpus_routes, mining

router = APIRouter()
router.include_router(chats.router)
router.include_router(sessions.router)
router.include_router(drafts.router)
router.include_router(corpus_routes.router)
router.include_router(mining.router)


# ═══════════════════════════════════════════════════════════════
#  Settings
# ═══════════════════════════════════════════════════════════════

class SearchSettingsRequest(BaseModel):
    api_key: str = ""
    provider: str = "duckduckgo"  # "duckduckgo" | "perplexity" | "brave" | "google"


@router.post("/settings/search")
async def update_search_settings(req: SearchSettingsRequest):
    from ..services import web_search
    key = req.api_key.strip() if req.api_key else None
    web_search.configure(api_key=key, provider=req.provider)
    return {"status": "ok", "provider": req.provider, "has_key": key is not None}


@router.get("/settings/search")
async def get_search_settings():
    from ..services import web_search
    return {
        "provider": web_search._api_provider,
        "has_key": web_search._api_key is not None,
    }


# ═══════════════════════════════════════════════════════════════
#  Health
# ═══════════════════════════════════════════════════════════════

@router.get("/health")
async def health():
    ollama_health = await ollama.health_check()
    corpus_errors = deps.corpus.validate()
    return {
        "status": "ok" if ollama_health["chat_model"] else "degraded",
        "ollama": ollama_health,
        "corpus_valid": len(corpus_errors) == 0,
        "corpus_errors": corpus_errors[:5],
    }


# ═══════════════════════════════════════════════════════════════
#  Model management
# ═══════════════════════════════════════════════════════════════

def _friendly_model_label(tag: str) -> str:
    """Turn a raw Ollama tag into a human-legible dropdown label.

    Examples:
      VladimirGav/gemma4-26b-16GB-VRAM:latest -> Gemma 4 26B
      gemma3:12b                              -> Gemma 3 12B
      gpt-oss:20b                             -> GPT-OSS 20B
      qwen3-vl:30b                            -> Qwen3 VL 30B
      deepseek-r1:14b                         -> DeepSeek R1 14B
      llama3.2:latest                         -> Llama 3.2
    """
    import re
    # Strip namespace, keep base + tag-suffix so we can search both for size.
    without_ns = tag.split("/")[-1].lower()
    base, _, tagsuffix = without_ns.partition(":")
    # Pull an explicit size token from either segment (base carries it for
    # community tags like gemma4-26b; tagsuffix carries it for canonical
    # Ollama tags like gemma3:12b).
    size_match = re.search(r"(\d+)b\b", base) or re.search(r"(\d+)b\b", tagsuffix)
    size = size_match.group(0).upper() if size_match else ""
    # Drop size + any noisy vram/ram qualifiers out of the stem
    stem = re.sub(r"-?\d+b\b", "", base)
    stem = re.sub(r"-?\d+gb[-_]?(v?ram)?", "", stem)
    stem = stem.strip("-_ ")
    # Canonicalize common family names
    family_map = {
        "gemma4": "Gemma 4",
        "gemma3": "Gemma 3",
        "gemma2": "Gemma 2",
        "gpt-oss": "GPT-OSS",
        "qwen3-vl": "Qwen3 VL",
        "qwen3": "Qwen3",
        "deepseek-r1": "DeepSeek R1",
        "llama3.2": "Llama 3.2",
        "llama3.1": "Llama 3.1",
    }
    pretty = family_map.get(stem, stem.replace("-", " ").title())
    return f"{pretty} {size}".strip() if size else pretty


# Models that are embedders / rerankers, not chat models. Filtered out of
# the chat dropdown so users can't accidentally pin nomic-embed-text as their
# conversational model.
_NON_CHAT_FAMILIES = ("embed", "rerank", "bge", "e5-")


@router.get("/models")
async def list_models():
    """List all locally available Ollama chat models + which one is active.

    Embedding / rerank models are filtered out — they're not valid chat
    targets. Each entry carries a `label` (human-friendly) and the raw
    `name` (what gets sent to /models/switch).
    """
    import httpx
    try:
        async with httpx.AsyncClient(base_url="http://localhost:11434") as client:
            resp = await client.get("/api/tags")
            resp.raise_for_status()
            models = resp.json().get("models", [])
            model_list = []
            for m in models:
                name = m["name"]
                if any(tok in name.lower() for tok in _NON_CHAT_FAMILIES):
                    continue
                model_list.append({
                    "name": name,
                    "label": _friendly_model_label(name),
                    "size": m.get("size", 0),
                })
            # Sort by label for stable UX
            model_list.sort(key=lambda x: x["label"].lower())
            return {"models": model_list, "active": ollama.CHAT_MODEL}
    except Exception as e:
        return {"models": [], "active": ollama.CHAT_MODEL, "error": str(e)}


class SwitchModelRequest(BaseModel):
    model: str


@router.post("/models/switch")
async def switch_model(req: SwitchModelRequest):
    """Hot-swap the active chat model.

    Unloads the current model from VRAM, loads the new one with
    keep_alive=-1 (pinned), and updates the global CHAT_MODEL.
    """
    import httpx
    old_model = ollama.CHAT_MODEL
    new_model = req.model

    try:
        async with httpx.AsyncClient(base_url="http://localhost:11434",
                                      timeout=httpx.Timeout(120.0)) as client:
            # Unload old model
            if old_model != new_model:
                await client.post("/api/generate", json={
                    "model": old_model, "prompt": "", "keep_alive": 0,
                })

            # Load new model with max GPU + pinned
            resp = await client.post("/api/generate", json={
                "model": new_model,
                "prompt": "",
                "stream": False,
                "keep_alive": -1,
                "options": {"num_gpu": 99, "num_ctx": ollama._NUM_CTX},
            })
            resp.raise_for_status()

            # Update the global model reference + auto-detect capabilities
            ollama.CHAT_MODEL = new_model
            from ..services import model_profiles
            profile = await model_profiles.set_active(new_model)

            # Inject model-switch boundary marker into active chat history
            # so the new model knows prior identity-adjacent nodes are historical
            if old_model != new_model:
                try:
                    # Find the most recent active chat
                    all_chats = chat_store.list_chats()
                    active = [c for c in all_chats if c.get("status") == "active"]
                    if active:
                        active_chat_id = active[0]["id"]
                        from ..models.schemas import ChatMessage
                        boundary = ChatMessage(
                            role="system",
                            content=(
                                f"[MODEL SWITCH] {old_model} → {new_model}. "
                                f"Prior messages were generated by {old_model}. "
                                f"Identity-related frame nodes from before this point are historical context, "
                                f"not your identity. You are {new_model} ({profile.family} family)."
                            ),
                        )
                        chat_store.append_message(active_chat_id, boundary)
                except Exception as e:
                    logger.warning("Failed to inject model-switch boundary: %s", e)

            # Build profile info for frontend
            profile_info = {
                "family": profile.family if profile else "",
                "params": profile.parameter_size if profile else "",
                "think": profile.supports_think if profile else False,
                "tools": profile.supports_tools if profile else False,
                "vision": profile.supports_vision if profile else False,
            }

            # Check actual VRAM usage
            ps_resp = await client.get("/api/ps")
            ps_data = ps_resp.json()
            for m in ps_data.get("models", []):
                if m["name"] == new_model:
                    return {
                        "ok": True,
                        "model": new_model,
                        "profile": profile_info,
                        "size_vram": m.get("size_vram", 0),
                        "size": m.get("size", 0),
                        "gpu_pct": round(
                            m.get("size_vram", 0) / max(1, m.get("size", 1)) * 100
                        ),
                    }
            return {"ok": True, "model": new_model, "profile": profile_info, "size_vram": 0, "size": 0, "gpu_pct": 0}
    except Exception as e:
        # Rollback on failure
        ollama.CHAT_MODEL = old_model
        return {"ok": False, "error": str(e), "model": old_model}


# ═══════════════════════════════════════════════════════════════
#  File upload
# ═══════════════════════════════════════════════════════════════

# Supported text-extractable file types
_TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".tsv", ".json", ".yaml", ".yml",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".html", ".css",
    ".c", ".cpp", ".h", ".hpp", ".java", ".go", ".rs", ".rb",
    ".sh", ".bash", ".zsh", ".ps1", ".bat",
    ".toml", ".ini", ".cfg", ".conf", ".env", ".xml",
    ".sql", ".r", ".m", ".swift", ".kt", ".scala", ".lua",
    ".log", ".tex",
}
_MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB


@router.post("/upload")
async def upload_file(file: UploadFile = File(...)):
    """Extract text content from an uploaded file.

    Supports plaintext/code files directly and PDFs via pdfplumber.
    Returns extracted text for the frontend to prepend to the user message.
    """
    import os
    ext = os.path.splitext(file.filename or "")[1].lower()

    # Read file bytes (with size guard)
    data = await file.read()
    if len(data) > _MAX_FILE_SIZE:
        raise HTTPException(413, f"File too large ({len(data)} bytes). Max is {_MAX_FILE_SIZE}.")

    extracted = ""

    if ext == ".pdf":
        try:
            import pdfplumber
            import io
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                pages = []
                for i, page in enumerate(pdf.pages):
                    text = page.extract_text() or ""
                    if text.strip():
                        pages.append(f"--- Page {i+1} ---\n{text}")
                extracted = "\n\n".join(pages)
                if not extracted.strip():
                    raise HTTPException(422, "PDF appears to contain no extractable text (scanned/image PDF).")
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, f"Failed to extract PDF text: {exc}")

    elif ext in (".docx", ".doc"):
        try:
            import docx
            import io
            doc = docx.Document(io.BytesIO(data))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        paragraphs.append(" | ".join(cells))
            extracted = "\n\n".join(paragraphs)
            if not extracted.strip():
                raise HTTPException(422, "DOCX appears to contain no extractable text.")
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(422, f"Failed to extract DOCX text: {exc}")

    elif ext in _TEXT_EXTENSIONS or ext == "":
        # Try decoding as text
        for encoding in ("utf-8", "utf-8-sig", "latin-1"):
            try:
                extracted = data.decode(encoding)
                break
            except (UnicodeDecodeError, ValueError):
                continue
        else:
            raise HTTPException(422, f"Could not decode file as text.")

    else:
        raise HTTPException(
            415,
            f"Unsupported file type '{ext}'. Supported: PDF, DOCX, and text/code files "
            f"({', '.join(sorted(list(_TEXT_EXTENSIONS)[:12]))}...)"
        )

    # Truncate very long files to avoid blowing up context
    char_limit = 80_000  # ~20k tokens
    truncated = False
    if len(extracted) > char_limit:
        extracted = extracted[:char_limit]
        truncated = True

    return {
        "filename": file.filename,
        "extension": ext,
        "chars": len(extracted),
        "truncated": truncated,
        "content": extracted,
    }


# ═══════════════════════════════════════════════════════════════
#  Events
# ═══════════════════════════════════════════════════════════════

@router.get("/events/{log_name}")
async def get_events(log_name: str, last_n: int = 50):
    valid_logs = [
        "gate_events.jsonl", "push_events.jsonl", "push_resolutions.jsonl",
        "degradation_flags.jsonl",
        "verification_log.jsonl", "drift_events.jsonl",
        "match_events.jsonl", "frame_events.jsonl", "proposal_events.jsonl",
    ]
    if log_name not in valid_logs:
        raise HTTPException(400, f"Invalid log name. Valid: {valid_logs}")
    return event_log.read_recent(log_name, last_n)
