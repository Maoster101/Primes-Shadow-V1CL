"""Repair mining sidecar source paths from a structured DOCX.

Uses stored evidence quotes first and explicit section numbers second. The command is
read-only unless --apply is supplied.
"""
from __future__ import annotations

import argparse
from difflib import SequenceMatcher
import json
import os
from pathlib import Path
import re
import unicodedata

from docx import Document


_MECHANICAL = re.compile(
    r"^(?:(?:segment|chunk|page|paragraph)\s*[-_ ]?\s*\d+|\d+(?:\.\d+)*|part\s+[ivxlcdm]+)$",
    re.IGNORECASE,
)
_SECTION_NUMBER = re.compile(r"(?<![A-Za-z0-9])((?:\d+|[A-Z])\.\d+)(?![A-Za-z0-9])")


def _norm(value: str) -> str:
    value = unicodedata.normalize("NFKD", value or "").casefold()
    value = re.sub(r"[^a-z0-9]+", " ", value)
    return re.sub(r"\s+", " ", value).strip()


def _document_index(path: Path):
    headings: dict[int, str] = {}
    paragraphs: list[tuple[str, list[str], str]] = []
    by_number: dict[str, list[str]] = {}
    for paragraph in Document(path).paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name if paragraph.style else "") or ""
        match = re.fullmatch(r"Heading\s+([1-6])", style, re.IGNORECASE)
        if match:
            level = int(match.group(1))
            for old_level in [key for key in headings if key >= level]:
                del headings[old_level]
            headings[level] = text
            source_path = [headings[key] for key in sorted(headings)]
            number = _SECTION_NUMBER.search(text)
            if number:
                by_number[number.group(1).casefold()] = source_path
            continue
        paragraphs.append((text, [headings[key] for key in sorted(headings)], _norm(text)))
    return paragraphs, by_number


def _path_from_sidecar(data: dict, paragraphs, by_number):
    quotes = [
        ref.get("quote", "") for ref in data.get("source_refs", [])
        if isinstance(ref, dict) and ref.get("quote")
    ]
    best = None
    for quote in quotes:
        normalized_quote = _norm(quote)
        for text, source_path, normalized_text in paragraphs:
            if len(normalized_quote) >= 20 and (
                normalized_quote in normalized_text or normalized_text in normalized_quote
            ):
                overlap = min(len(normalized_quote), len(normalized_text)) / max(
                    len(normalized_quote), len(normalized_text)
                )
                candidate = (2.0 + overlap, source_path)
                if best is None or candidate[0] > best[0]:
                    best = candidate
    if best:
        return best[1], "quote"

    for value in [
        *(data.get("source_topic") or []),
        *(ref.get("locator", "") for ref in data.get("source_refs", []) if isinstance(ref, dict)),
    ]:
        for number in _SECTION_NUMBER.findall(str(value)):
            if number.casefold() in by_number:
                return by_number[number.casefold()], "section_number"

    # Fuzzy matching is deliberately last and conservative.
    for quote in quotes:
        normalized_quote = _norm(quote)
        for text, source_path, normalized_text in paragraphs:
            score = SequenceMatcher(None, normalized_quote, normalized_text).ratio()
            if best is None or score > best[0]:
                best = (score, source_path)
    if best and best[0] >= 0.62:
        return best[1], "fuzzy_quote"

    current = data.get("source_topic") or data.get("source_path") or []
    if isinstance(current, str):
        current = [part.strip() for part in current.split("/") if part.strip()]
    descriptive = [part for part in current if not _MECHANICAL.fullmatch(str(part).strip())]
    return descriptive, "retained" if descriptive else "unresolved"


def _source_position(data: dict, source_path: list[str], document_text: str) -> int:
    """Locate evidence in original document text; headings are the safe fallback."""
    folded = document_text.casefold()
    positions: list[int] = []
    for ref in data.get("source_refs") or []:
        if not isinstance(ref, dict):
            continue
        quote = str(ref.get("quote") or "").strip().casefold()
        for fragment in [part.strip(" .…") for part in re.split(r"\.{3}|…", quote)]:
            if len(fragment) < 16:
                continue
            position = folded.find(fragment)
            if position >= 0:
                positions.append(position)
                break
    if positions:
        return min(positions)
    for heading in reversed(source_path):
        position = folded.find(str(heading).casefold())
        if position >= 0:
            return position
    return 10**12


def _temp_suffix(data: dict) -> int:
    match = re.search(r"(\d+)$", str(data.get("temp_id") or ""))
    return int(match.group(1)) if match else 0

def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--session-dir", type=Path, required=True)
    parser.add_argument("--docx", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()

    paragraphs, by_number = _document_index(args.docx)
    document_text = "\n\n".join(
        paragraph.text.strip() for paragraph in Document(args.docx).paragraphs
        if paragraph.text.strip()
    )
    counts: dict[str, int] = {}
    unresolved: list[str] = []
    changed = 0
    for path in sorted(args.session_dir.glob("*_raw.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        if data.get("type") != "slab":
            continue
        source_path, method = _path_from_sidecar(data, paragraphs, by_number)
        counts[method] = counts.get(method, 0) + 1
        if not source_path:
            unresolved.append(path.name)
            continue
        old = data.get("source_topic") or data.get("source_path") or []
        source_order = _source_position(data, source_path, document_text) * 1_000 + _temp_suffix(data)
        needs_update = (
            old != source_path or data.get("source_path") != source_path
            or data.get("_source_order") != source_order
        )
        if needs_update:
            changed += 1
            if args.apply:
                data["source_path"] = source_path
                data["source_topic"] = source_path
                data["_source_order"] = source_order
                temp = path.with_suffix(path.suffix + ".tmp")
                temp.write_text(
                    json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
                )
                os.replace(temp, path)

    print(json.dumps({
        "mode": "apply" if args.apply else "dry-run",
        "changed": changed,
        "methods": counts,
        "unresolved": unresolved,
    }, ensure_ascii=False, indent=2))
    return 1 if unresolved else 0


if __name__ == "__main__":
    raise SystemExit(main())